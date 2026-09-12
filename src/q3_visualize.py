"""绘制问题 3 长时干燥过程的论文用可视化图。

当前仓库尚未保存 Q3 的未四舍五入结构化求解结果，因此本脚本只读
``outputs/q3/workbooks/result3.xlsx`` 中的交付表数据。工作簿数值保留四位
小数，适合展示演化趋势和空间分布，但不能单独验证严格不等式
``max(C) < 0.15``。正式定稿时应将数据源替换为 Kirchhoff 界面离散求解器
导出的 NPZ/NPY 原始结果。
"""

from __future__ import annotations

import argparse
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path

matplotlib_cache = Path(tempfile.gettempdir()) / "cumcm2026_matplotlib"
matplotlib_cache.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(matplotlib_cache))

import matplotlib
import numpy as np

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.cm import ScalarMappable
from matplotlib.colors import LinearSegmentedColormap, LogNorm, Normalize
from openpyxl import load_workbook

from visualization_style import (
    COLORS,
    RADIAL_COLORS,
    RADIAL_STYLES,
    add_panel_label,
    configure_publication_style,
    polish_axis,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "outputs" / "q3" / "workbooks" / "result3.xlsx"
DEFAULT_FIGURE_DIR = PROJECT_ROOT / "outputs" / "q3" / "figures" / "fancy"
TARGET_MOISTURE = 0.15
SELECTED_RADII_CM = (0.0, 0.5, 1.0, 1.5, 2.0)
PROFILE_TIMES_H = (6.0, 12.0, 24.0, 36.0, 48.0, 54.0, 57.5)
CROSS_SECTION_TIMES_H = (6.0, 12.0, 24.0, 36.0, 48.0, 57.5)
MOISTURE_CMAP = "YlGnBu"
TIME_CMAP = LinearSegmentedColormap.from_list(
    "cumcm_q3_time",
    (COLORS["navy"], COLORS["blue"], COLORS["teal"], COLORS["gold"], COLORS["vermilion"]),
)


@dataclass(frozen=True)
class Q3Data:
    """Q3 交付工作簿中的径向含水率历史。"""

    time_s: np.ndarray
    radius_cm: np.ndarray
    moisture: np.ndarray
    sheet_name: str
    source_range: str


def load_q3_workbook(path: Path) -> Q3Data:
    """只读加载并校验 Q3 工作簿。"""

    if not path.is_file():
        raise FileNotFoundError(f"未找到 Q3 工作簿：{path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("Q3 工作簿应只包含一个数据工作表")
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if len(rows) < 2 or len(rows[0]) < 3:
        raise ValueError("Q3 工作簿没有足够的数据行或径向节点")
    if any(value is None for row in rows for value in row):
        raise ValueError("Q3 工作簿的数据区域存在空单元格")

    radius_cm = np.asarray(rows[0][1:], dtype=float)
    numerical = np.asarray(rows[1:], dtype=float)
    time_s = numerical[:, 0]
    moisture = numerical[:, 1:]
    expected_shape = (time_s.size, radius_cm.size)
    if moisture.shape != expected_shape:
        raise ValueError(
            f"含水率矩阵尺寸错误：期望 {expected_shape}，实际 {moisture.shape}"
        )
    if not np.all(np.diff(time_s) > 0) or not np.all(np.diff(radius_cm) > 0):
        raise ValueError("时间和半径坐标必须严格递增")
    if not np.isfinite(moisture).all() or np.any(moisture <= 0):
        raise ValueError("含水率中存在非有限值或非正值")
    if np.any(np.diff(moisture, axis=0) > 5e-12):
        raise ValueError("检测到含水率随时间反弹，需先核查源数据")
    if np.any(np.diff(moisture, axis=1) > 5e-12):
        raise ValueError("检测到含水率沿半径反向增大，需先核查源数据")

    last_column = _excel_column_name(len(rows[0]))
    return Q3Data(
        time_s=time_s,
        radius_cm=radius_cm,
        moisture=moisture,
        sheet_name=sheet.title,
        source_range=f"A1:{last_column}{len(rows)}",
    )


def _excel_column_name(index: int) -> str:
    """将从 1 开始的列号转换为 Excel 列名。"""

    letters = []
    while index:
        index, remainder = divmod(index - 1, 26)
        letters.append(chr(65 + remainder))
    return "".join(reversed(letters))


def nearest_indices(values: np.ndarray, targets: tuple[float, ...]) -> list[int]:
    """返回目标坐标的最近下标。"""

    if min(targets) < values[0] or max(targets) > values[-1]:
        raise ValueError("指定的绘图位置或时刻超出结果范围")
    return [int(np.argmin(np.abs(values - target))) for target in targets]


def radial_label(radius_cm: float, maximum_radius_cm: float) -> str:
    if np.isclose(radius_cm, 0.0):
        return "中心"
    if np.isclose(radius_cm, maximum_radius_cm):
        return "表面"
    return f"r = {radius_cm:g} cm"


def radial_style(radius_cm: float, maximum_radius_cm: float) -> tuple[str, object]:
    slot = int(round(radius_cm / maximum_radius_cm * (len(RADIAL_COLORS) - 1)))
    slot = min(max(slot, 0), len(RADIAL_COLORS) - 1)
    return RADIAL_COLORS[slot], RADIAL_STYLES[slot]


def save_figure(
    figure: plt.Figure,
    output_directory: Path,
    stem: str,
    formats: tuple[str, ...],
) -> list[Path]:
    output_directory.mkdir(parents=True, exist_ok=True)
    paths = []
    for suffix in formats:
        path = output_directory / f"{stem}.{suffix}"
        figure.savefig(
            path,
            dpi=300,
            bbox_inches="tight",
            metadata={"Creator": "CUMCM2026 visualization pipeline"},
        )
        paths.append(path.resolve())
    plt.close(figure)
    return paths


def rounded_completion_times_h(data: Q3Data) -> np.ndarray:
    """由四位小数展示值计算首次达到 C≤0.1500 的时刻。"""

    completion = np.full(data.radius_cm.size, np.nan)
    for column in range(data.radius_cm.size):
        indices = np.flatnonzero(data.moisture[:, column] <= TARGET_MOISTURE)
        if indices.size:
            completion[column] = data.time_s[indices[0]] / 3600.0
    return completion


def draw_selected_histories(axis, data: Q3Data, *, legend: bool = True) -> None:
    time_h = data.time_s / 3600.0
    indices = nearest_indices(data.radius_cm, SELECTED_RADII_CM)
    for index in indices:
        color, style = radial_style(data.radius_cm[index], data.radius_cm[-1])
        axis.plot(
            time_h,
            data.moisture[:, index],
            color=color,
            linestyle=style,
            linewidth=2.1,
            label=radial_label(data.radius_cm[index], data.radius_cm[-1]),
        )
    axis.axhline(
        TARGET_MOISTURE,
        color=COLORS["vermilion"],
        linestyle=(0, (3, 2)),
        linewidth=1.35,
        label="目标值 0.15",
    )
    axis.set(
        xlabel="干燥时间/h",
        ylabel="干基含水率 C/(kg·kg⁻¹)",
        xlim=(0.0, time_h[-1]),
    )
    polish_axis(axis, grid_axis="y")
    if legend:
        axis.legend(ncol=3, loc="upper right", handlelength=2.8, columnspacing=1.25)


def plot_threshold_judgement(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制全域最大含水率与终止阈值，并放大末段。"""

    time_h = data.time_s / 3600.0
    maximum = np.max(data.moisture, axis=1)
    displayed_hits = np.flatnonzero(maximum <= TARGET_MOISTURE)
    displayed_time = time_h[displayed_hits[0]] if displayed_hits.size else np.nan

    figure, axis = plt.subplots(figsize=(10.7, 5.9), constrained_layout=True)
    axis.plot(time_h, maximum, color=COLORS["navy"], linewidth=2.5, label="全域最大值")
    axis.fill_between(time_h, TARGET_MOISTURE, maximum, color=COLORS["blue"], alpha=0.10)
    axis.axhline(
        TARGET_MOISTURE,
        color=COLORS["vermilion"],
        linestyle=(0, (4, 2)),
        linewidth=1.7,
        label="终止阈值 C = 0.15",
    )
    axis.axvline(time_h[-1], color=COLORS["gold"], linewidth=1.5, linestyle=(0, (2, 2)))
    axis.scatter(
        [time_h[-1]],
        [maximum[-1]],
        s=70,
        color=COLORS["vermilion"],
        edgecolor="white",
        linewidth=1.2,
        zorder=5,
    )
    axis.annotate(
        f"交付表末行\n{time_h[-1]:.2f} h，显示值 {maximum[-1]:.4f}",
        xy=(time_h[-1], maximum[-1]),
        xytext=(38.5, 0.59),
        arrowprops={"arrowstyle": "->", "color": COLORS["vermilion"], "lw": 1.2},
        color=COLORS["vermilion"],
        fontsize=10.2,
        ha="left",
    )
    axis.set(
        title="Q3 终止判据：最慢位置决定总干燥时间",
        xlabel="干燥时间/h",
        ylabel="全域最大干基含水率 Cmax/(kg·kg⁻¹)",
        xlim=(0.0, time_h[-1] + 0.7),
        ylim=(0.08, 2.68),
    )
    polish_axis(axis, grid_axis="y")
    axis.legend(loc="upper right")

    inset = axis.inset_axes([0.48, 0.40, 0.47, 0.42])
    inset.plot(time_h, maximum, color=COLORS["navy"], linewidth=2.0)
    inset.axhline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle="--", linewidth=1.2)
    if np.isfinite(displayed_time):
        inset.axvline(displayed_time, color=COLORS["gold"], linewidth=1.25)
        inset.scatter([displayed_time], [TARGET_MOISTURE], s=33, color=COLORS["gold"], zorder=4)
    inset.set(xlim=(48.0, time_h[-1] + 0.12), ylim=(0.1485, 0.1635), title="末段放大")
    inset.tick_params(labelsize=8.2, direction="out", length=3)
    inset.grid(True, color=COLORS["light_gray"], linewidth=0.55, alpha=0.75)
    for spine in ("top", "right"):
        inset.spines[spine].set_visible(False)

    axis.text(
        0.0,
        -0.18,
        "注：工作簿仅保留四位小数，末行 0.1500 不能单独证明严格不等式 Cmax<0.15；正式判定以未四舍五入结果为准。",
        transform=axis.transAxes,
        ha="left",
        va="top",
        fontsize=8.9,
        color=COLORS["gray"],
        clip_on=False,
    )
    return save_figure(figure, output_directory, "q3_threshold_judgement", formats)


def plot_radial_histories(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制多半径含水率历史及末段放大图。"""

    time_h = data.time_s / 3600.0
    figure, axes = plt.subplots(1, 2, figsize=(12.4, 5.1), constrained_layout=True)
    draw_selected_histories(axes[0], data, legend=True)
    axes[0].set(title="全程响应：表面先干，中心后达标", ylim=(0.0, 2.68))
    add_panel_label(axes[0], "a")

    indices = nearest_indices(data.radius_cm, SELECTED_RADII_CM)
    for index in indices:
        color, style = radial_style(data.radius_cm[index], data.radius_cm[-1])
        axes[1].plot(
            time_h,
            data.moisture[:, index],
            color=color,
            linestyle=style,
            linewidth=2.1,
        )
    axes[1].axhline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle="--", linewidth=1.4)
    axes[1].fill_between(
        [40.0, time_h[-1]],
        [0.0, 0.0],
        [TARGET_MOISTURE, TARGET_MOISTURE],
        color=COLORS["teal"],
        alpha=0.08,
    )
    axes[1].scatter(
        [time_h[-1], time_h[-1]],
        [data.moisture[-1, 0], data.moisture[-1, -1]],
        color=(COLORS["navy"], COLORS["vermilion"]),
        s=48,
        zorder=4,
    )
    axes[1].annotate(
        f"中心 {data.moisture[-1, 0]:.4f}",
        xy=(time_h[-1], data.moisture[-1, 0]),
        xytext=(50.2, 0.205),
        arrowprops={"arrowstyle": "->", "color": COLORS["navy"], "lw": 1.0},
        color=COLORS["navy"],
    )
    axes[1].annotate(
        f"表面 {data.moisture[-1, -1]:.4f}",
        xy=(time_h[-1], data.moisture[-1, -1]),
        xytext=(46.2, 0.078),
        arrowprops={"arrowstyle": "->", "color": COLORS["vermilion"], "lw": 1.0},
        color=COLORS["vermilion"],
    )
    axes[1].set(
        title="末段放大：中心控制终止时刻",
        xlabel="干燥时间/h",
        ylabel="干基含水率 C/(kg·kg⁻¹)",
        xlim=(40.0, time_h[-1] + 0.4),
        ylim=(0.035, 0.21),
    )
    polish_axis(axes[1], grid_axis="both")
    add_panel_label(axes[1], "b")
    return save_figure(figure, output_directory, "q3_radial_histories", formats)


def draw_heatmap(axis, data: Q3Data, *, add_colorbar: bool = True):
    time_h = data.time_s / 3600.0
    norm = LogNorm(vmin=float(np.min(data.moisture)), vmax=float(np.max(data.moisture)))
    mesh = axis.pcolormesh(
        time_h,
        data.radius_cm,
        data.moisture.T,
        shading="auto",
        cmap=MOISTURE_CMAP,
        norm=norm,
        rasterized=True,
    )
    contour = axis.contour(
        time_h,
        data.radius_cm,
        data.moisture.T,
        levels=[TARGET_MOISTURE],
        colors=[COLORS["vermilion"]],
        linewidths=1.8,
    )
    if contour.allsegs and contour.allsegs[0]:
        axis.clabel(contour, fmt={TARGET_MOISTURE: "C = 0.15"}, fontsize=9, inline=True)
    axis.set(
        xlabel="干燥时间/h",
        ylabel="径向位置 r/cm",
        xlim=(0.0, time_h[-1]),
        ylim=(data.radius_cm[0], data.radius_cm[-1]),
    )
    axis.set_yticks([0.0, 0.5, 1.0, 1.5, 2.0], ["中心", "0.5", "1.0", "1.5", "表面"])
    if add_colorbar:
        colorbar = axis.figure.colorbar(mesh, ax=axis, pad=0.025, fraction=0.045)
        colorbar.set_label("干基含水率 C/(kg·kg⁻¹)，对数色标")
        colorbar.set_ticks([0.05, 0.10, 0.15, 0.30, 0.60, 1.20, 2.40])
        colorbar.set_ticklabels(["0.05", "0.10", "0.15", "0.30", "0.60", "1.20", "2.40"])
    return mesh


def plot_spatiotemporal_map(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制长时干燥含水率时空热图。"""

    figure, axis = plt.subplots(figsize=(11.2, 5.2), constrained_layout=True)
    draw_heatmap(axis, data)
    axis.set_title("Q3 含水率时空图：达标等值线由表面向中心推进")
    axis.annotate(
        "表面约 12.6 h\n达到显示阈值",
        xy=(12.58, data.radius_cm[-1]),
        xytext=(19.0, 1.66),
        arrowprops={"arrowstyle": "->", "color": COLORS["vermilion"], "lw": 1.1},
        color=COLORS["vermilion"],
        fontsize=9.4,
    )
    axis.text(
        0.012,
        0.035,
        "颜色表示含水率大小；深色含水率高，浅色含水率低",
        transform=axis.transAxes,
        fontsize=9.2,
        color=COLORS["charcoal"],
        bbox={"boxstyle": "round,pad=0.3", "facecolor": "white", "edgecolor": "none", "alpha": 0.82},
    )
    return save_figure(figure, output_directory, "q3_moisture_spacetime", formats)


def plot_radial_profiles(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制多个代表时刻的径向含水率剖面。"""

    target_s = tuple(hour * 3600.0 for hour in PROFILE_TIMES_H)
    indices = nearest_indices(data.time_s, target_s)
    colors = TIME_CMAP(np.linspace(0.03, 0.97, len(indices)))
    figure, axis = plt.subplots(figsize=(9.7, 5.9), constrained_layout=True)
    for index, color in zip(indices, colors, strict=True):
        axis.plot(
            data.radius_cm,
            data.moisture[index],
            color=color,
            linewidth=2.15,
            marker="o" if index == indices[-1] else None,
            markersize=3.8,
            label=f"{data.time_s[index] / 3600:g} h",
        )
    axis.axhline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle="--", linewidth=1.4)
    axis.fill_between(
        data.radius_cm,
        0.0,
        TARGET_MOISTURE,
        color=COLORS["teal"],
        alpha=0.065,
    )
    axis.annotate(
        "末时刻仍由中心值控制",
        xy=(0.0, data.moisture[-1, 0]),
        xytext=(0.42, 0.39),
        arrowprops={"arrowstyle": "->", "color": COLORS["navy"], "lw": 1.1},
        color=COLORS["navy"],
    )
    axis.set(
        title="径向剖面：内部水分梯度在长时干燥中持续存在",
        xlabel="径向位置 r/cm",
        ylabel="干基含水率 C/(kg·kg⁻¹)",
        xlim=(data.radius_cm[0], data.radius_cm[-1]),
        ylim=(0.0, 1.08),
    )
    polish_axis(axis, grid_axis="y")
    axis.legend(title="干燥时间", ncol=2, loc="upper right")
    return save_figure(figure, output_directory, "q3_radial_profiles", formats)


def plot_cross_section_evolution(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """以六个圆截面展示含水率空间分布随时间的变化。"""

    target_s = tuple(hour * 3600.0 for hour in CROSS_SECTION_TIMES_H)
    indices = nearest_indices(data.time_s, target_s)
    selected = data.moisture[indices]
    norm = LogNorm(vmin=float(np.min(selected)), vmax=float(np.max(selected)))
    cmap = plt.get_cmap(MOISTURE_CMAP).copy()
    cmap.set_bad("white", alpha=0.0)
    radius = data.radius_cm[-1]
    coordinates = np.linspace(-radius, radius, 301)
    xx, yy = np.meshgrid(coordinates, coordinates)
    radial_distance = np.sqrt(xx**2 + yy**2)

    figure, axes = plt.subplots(2, 3, figsize=(11.5, 7.65), constrained_layout=True)
    for axis, index in zip(axes.ravel(), indices, strict=True):
        field = np.interp(radial_distance, data.radius_cm, data.moisture[index])
        field = np.ma.masked_where(radial_distance > radius, field)
        image = axis.imshow(
            field,
            extent=(-radius, radius, -radius, radius),
            origin="lower",
            cmap=cmap,
            norm=norm,
            interpolation="bilinear",
        )
        axis.add_patch(
            plt.Circle((0.0, 0.0), radius, fill=False, color=COLORS["charcoal"], linewidth=0.8)
        )
        axis.scatter([0.0], [0.0], s=7, color="white", alpha=0.85)
        axis.set(
            title=(
                f"{data.time_s[index] / 3600:g} h\n"
                f"中心 {data.moisture[index, 0]:.3f} · 表面 {data.moisture[index, -1]:.3f}"
            ),
            xlim=(-radius - 0.08, radius + 0.08),
            ylim=(-radius - 0.08, radius + 0.08),
            aspect="equal",
        )
        axis.axis("off")

    colorbar = figure.colorbar(image, ax=axes.ravel().tolist(), pad=0.025, fraction=0.035)
    colorbar.set_label("干基含水率 C/(kg·kg⁻¹)")
    colorbar.set_ticks([0.05, 0.10, 0.15, 0.30, 0.60, 1.00])
    colorbar.set_ticklabels(["0.05", "0.10", "0.15", "0.30", "0.60", "1.00"])
    figure.suptitle("Q3 圆截面含水率演化：低含水区从外层逐步向中心扩展", fontsize=14)
    figure.text(
        0.5,
        0.012,
        "同一色标用于全部截面；深色含水率高，浅色含水率低",
        ha="center",
        fontsize=9.2,
        color=COLORS["gray"],
    )
    return save_figure(figure, output_directory, "q3_cross_section_evolution", formats)


def plot_completion_front(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制各径向位置首次显示为 C≤0.1500 的时刻。"""

    completion_h = rounded_completion_times_h(data)
    valid = np.isfinite(completion_h)
    if not np.all(valid):
        raise ValueError("至少一个径向位置未在工作簿内达到显示阈值 C≤0.1500")

    figure, axis = plt.subplots(figsize=(9.5, 5.8), constrained_layout=True)
    axis.fill_between(
        data.radius_cm,
        completion_h,
        np.full_like(completion_h, completion_h.max() + 2.5),
        color=COLORS["blue"],
        alpha=0.09,
        label="尚未达到显示阈值",
    )
    axis.fill_between(
        data.radius_cm,
        np.zeros_like(completion_h),
        completion_h,
        color=COLORS["teal"],
        alpha=0.10,
        label="该位置已达到显示阈值",
    )
    axis.plot(data.radius_cm, completion_h, color=COLORS["navy"], linewidth=2.35)
    scatter = axis.scatter(
        data.radius_cm,
        completion_h,
        c=completion_h,
        cmap=TIME_CMAP,
        norm=Normalize(vmin=float(completion_h.min()), vmax=float(completion_h.max())),
        s=58,
        edgecolor="white",
        linewidth=0.8,
        zorder=4,
    )
    axis.annotate(
        f"表面最先\n{completion_h[-1]:.2f} h",
        xy=(data.radius_cm[-1], completion_h[-1]),
        xytext=(1.52, 20.5),
        arrowprops={"arrowstyle": "->", "color": COLORS["vermilion"], "lw": 1.2},
        color=COLORS["vermilion"],
        ha="center",
    )
    axis.annotate(
        f"中心最后\n{completion_h[0]:.2f} h",
        xy=(data.radius_cm[0], completion_h[0]),
        xytext=(0.30, 50.7),
        arrowprops={"arrowstyle": "->", "color": COLORS["navy"], "lw": 1.2},
        color=COLORS["navy"],
        ha="center",
    )
    axis.set(
        title="达标前沿：外层先干，判据逐层向中心推进",
        xlabel="径向位置 r/cm（0 为中心，2 为表面）",
        ylabel="首次显示为 C≤0.1500 的时间/h",
        xlim=(-0.04, data.radius_cm[-1] + 0.04),
        ylim=(0.0, completion_h.max() + 3.5),
    )
    polish_axis(axis, grid_axis="y")
    axis.legend(loc="lower left")
    colorbar = figure.colorbar(scatter, ax=axis, pad=0.025, fraction=0.045)
    colorbar.set_label("达到显示阈值的时间/h")
    axis.text(
        0.0,
        -0.18,
        "注：本图使用工作簿四位小数显示值的 C≤0.1500；严格判据仍需未四舍五入结果。",
        transform=axis.transAxes,
        ha="left",
        va="top",
        fontsize=8.8,
        color=COLORS["gray"],
        clip_on=False,
    )
    return save_figure(figure, output_directory, "q3_completion_front", formats)


def plot_center_surface_phase(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制中心—表面含水率相图，展示径向滞后。"""

    time_h = data.time_s / 3600.0
    center = data.moisture[:, 0]
    surface = data.moisture[:, -1]
    figure, axis = plt.subplots(figsize=(7.7, 6.8), constrained_layout=True)
    axis.plot(center, surface, color=COLORS["light_gray"], linewidth=3.2, zorder=1)
    sample = np.unique(np.linspace(0, time_h.size - 1, 420).astype(int))
    scatter = axis.scatter(
        center[sample],
        surface[sample],
        c=time_h[sample],
        cmap=TIME_CMAP,
        s=26,
        linewidth=0.0,
        zorder=3,
    )
    lower = min(float(surface.min()), TARGET_MOISTURE) * 0.88
    upper = float(center.max()) * 1.08
    axis.plot([lower, upper], [lower, upper], color=COLORS["gray"], linestyle="--", linewidth=1.1)
    axis.axvline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle=(0, (3, 2)), linewidth=1.1)
    axis.axhline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle=(0, (3, 2)), linewidth=1.1)
    axis.annotate(
        "表面响应领先\n轨迹落在等值线下方",
        xy=(0.46, 0.165),
        xytext=(0.79, 0.095),
        arrowprops={"arrowstyle": "->", "color": COLORS["teal"], "lw": 1.2},
        color=COLORS["teal"],
        fontsize=10,
    )
    axis.scatter(
        [center[-1]],
        [surface[-1]],
        s=78,
        color=COLORS["vermilion"],
        edgecolor="white",
        linewidth=1.1,
        zorder=5,
    )
    axis.set(
        title="中心—表面相图：一条轨迹读出空间滞后",
        xlabel="中心干基含水率 C(center)/(kg·kg⁻¹)",
        ylabel="表面干基含水率 C(surface)/(kg·kg⁻¹)",
        xscale="log",
        yscale="log",
        xlim=(lower, upper),
        ylim=(lower, upper),
        aspect="equal",
    )
    polish_axis(axis, grid_axis="both")
    colorbar = figure.colorbar(scatter, ax=axis, pad=0.025, fraction=0.05)
    colorbar.set_label("干燥时间/h")
    axis.text(
        1.25,
        1.45,
        "虚线：中心与表面含水率相同",
        color=COLORS["gray"],
        rotation=43,
        fontsize=8.8,
    )
    return save_figure(figure, output_directory, "q3_center_surface_phase", formats)


def plot_q3_overview(
    data: Q3Data, output_directory: Path, formats: tuple[str, ...]
) -> list[Path]:
    """绘制适合论文与答辩首页使用的四面板总览。"""

    time_h = data.time_s / 3600.0
    maximum = np.max(data.moisture, axis=1)
    figure, axes = plt.subplots(2, 2, figsize=(13.0, 9.0), constrained_layout=True)

    axes[0, 0].plot(time_h, maximum, color=COLORS["navy"], linewidth=2.35)
    axes[0, 0].axhline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle="--", linewidth=1.4)
    axes[0, 0].scatter([time_h[-1]], [maximum[-1]], s=62, color=COLORS["vermilion"], zorder=4)
    axes[0, 0].set(
        title="最慢位置控制总干燥时间",
        xlabel="干燥时间/h",
        ylabel="全域最大含水率 Cmax/(kg·kg⁻¹)",
        xlim=(0.0, time_h[-1] + 0.5),
        ylim=(0.08, 2.66),
    )
    polish_axis(axes[0, 0], grid_axis="y")
    axes[0, 0].text(
        0.97,
        0.68,
        f"交付表终点\n{time_h[-1]:.2f} h",
        transform=axes[0, 0].transAxes,
        ha="right",
        color=COLORS["vermilion"],
        fontsize=10.4,
    )
    add_panel_label(axes[0, 0], "a")

    draw_selected_histories(axes[0, 1], data, legend=False)
    axes[0, 1].set(title="径向响应不同步", ylim=(0.0, 2.66))
    axes[0, 1].legend(ncol=2, loc="upper right", handlelength=2.5, columnspacing=1.0)
    add_panel_label(axes[0, 1], "b")

    mesh = draw_heatmap(axes[1, 0], data, add_colorbar=False)
    axes[1, 0].set_title("含水率时空演化（对数色标）")
    add_panel_label(axes[1, 0], "c")
    colorbar = figure.colorbar(mesh, ax=axes[1, 0], pad=0.02, fraction=0.045)
    colorbar.set_label("C/(kg·kg⁻¹)")

    axes[1, 1].plot(
        data.radius_cm,
        data.moisture[-1],
        color=COLORS["navy"],
        linewidth=2.4,
        marker="o",
        markersize=4.0,
    )
    axes[1, 1].axhline(TARGET_MOISTURE, color=COLORS["vermilion"], linestyle="--", linewidth=1.4)
    axes[1, 1].fill_between(
        data.radius_cm,
        0.0,
        TARGET_MOISTURE,
        color=COLORS["teal"],
        alpha=0.08,
    )
    axes[1, 1].annotate(
        f"中心 {data.moisture[-1, 0]:.4f}",
        xy=(0.0, data.moisture[-1, 0]),
        xytext=(0.36, 0.163),
        arrowprops={"arrowstyle": "->", "color": COLORS["navy"], "lw": 1.0},
        color=COLORS["navy"],
    )
    axes[1, 1].annotate(
        f"表面 {data.moisture[-1, -1]:.4f}",
        xy=(2.0, data.moisture[-1, -1]),
        xytext=(1.36, 0.078),
        arrowprops={"arrowstyle": "->", "color": COLORS["vermilion"], "lw": 1.0},
        color=COLORS["vermilion"],
    )
    axes[1, 1].set(
        title="末时刻径向剖面",
        xlabel="径向位置 r/cm",
        ylabel="干基含水率 C/(kg·kg⁻¹)",
        xlim=(0.0, 2.0),
        ylim=(0.035, 0.165),
    )
    polish_axis(axes[1, 1], grid_axis="both")
    add_panel_label(axes[1, 1], "d")

    figure.suptitle("问题 3｜长时干燥过程与终止判据", fontsize=15.2)
    figure.text(
        0.5,
        0.005,
        "数据源为四位小数交付表；严格阈值判定须以未四舍五入结果复核",
        ha="center",
        fontsize=8.9,
        color=COLORS["gray"],
    )
    return save_figure(figure, output_directory, "q3_overview", formats)


def create_all_figures(
    data: Q3Data,
    output_directory: Path,
    formats: tuple[str, ...],
) -> list[Path]:
    paths = []
    for plotter in (
        plot_q3_overview,
        plot_threshold_judgement,
        plot_radial_histories,
        plot_spatiotemporal_map,
        plot_radial_profiles,
        plot_cross_section_evolution,
        plot_completion_front,
        plot_center_surface_phase,
    ):
        paths.extend(plotter(data, output_directory, formats))
    return paths


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="绘制问题 3 的论文用长时干燥图")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_FIGURE_DIR)
    parser.add_argument(
        "--formats",
        nargs="+",
        default=("png", "pdf", "svg"),
        choices=("png", "pdf", "svg"),
    )
    return parser.parse_args()


def main() -> None:
    arguments = parse_arguments()
    configure_publication_style()
    data = load_q3_workbook(arguments.workbook.resolve())
    paths = create_all_figures(
        data,
        arguments.output_dir.resolve(),
        tuple(arguments.formats),
    )
    print(
        f"Q3 图已生成：{len(paths)} 个文件；"
        f"数据区域 {data.sheet_name}!{data.source_range}"
    )
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()

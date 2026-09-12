"""
2026 高教社杯 A题 —— 问题4 求解器（收缩条件下的移动边界热-湿耦合）

模型口径见 docs/q4/moving_boundary_heat_moisture_derivation.md 与 PROJECT.md §4.4：
  * 材料坐标 xi = r/R(t) 把移动区域映射到固定区间 [0,1]。均匀径向相似收缩下，物理
    坐标中的材料对流项 u dphi/dr 与坐标移动项严格抵消，**固定域方程中不得再加任何
    收缩对流项**，否则重复计算材料运动；
  * 收缩只通过两条路径进入模型：内部传输算子的 R(t)^-2，表面 Robin 项的 R(t)^-1；
  * 物性统一用附录4 的 rho(C)、cp(C)、k(C)、D(C,T_K)，D 中温度必须用 K；不考虑潜热；
  * h、hm 沿用问题1 的题给值（题目未给问题4 的界面系数）；
  * R(t) 取附件2 分段线性插值，t>72h 保持末值并显式报告外推；
  * 环境边界与问题3 完全同一口径：附件1 内逐点插值，4h 后取 t>=9000s 平台均值延拓。

离散（与 q2/q3 同族，但界面系数必须换掉）：
  * 顶点型守恒有限体积，节点 xi_i = i/N，xi=0 与 xi=1 都落在节点上；
  * 水分界面系数用 Kirchhoff 积分平均（PROJECT.md §12 第7条）。**不复用 q2 的
    harm()**：点值调和平均在 D 退化区会让表面节点与内部完全解耦，制造"结壳停滞"
    数值假象并破坏网格收敛；
  * 导热系数 k 只跨 2.3 倍且不退化，界面仍取算术平均；
  * 后向欧拉 + 每个时间层内 Picard 交替更新 C -> 物性 -> T；三对角直接求解。

复用 q2_coupled 的 make_grid / assemble / bsolve / check_K / arith：这些只依赖网格
几何与守恒组装，与物性、与是否移动边界无关，在 xi 坐标上同样成立（见下方推导注释）。

用法：
    python src/q4_moving_boundary.py            # 全流程：校验、极限测试、收敛、正式解、导出
    python src/q4_moving_boundary.py --quick    # 只跑正式解，跳过收敛与对照实验
"""
import os
import shutil
import sys
import warnings

import numpy as np
import pandas as pd
from scipy.interpolate import interp1d
from scipy.special import exp1

from q2_coupled import (P, rho_f, cp_f, k_f, make_grid, assemble, bsolve,
                        arith, check_K)
from q3_drying_time import load_env_ext

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OFFICIAL = os.path.join(PROJECT_ROOT, 'data', 'official', 'A题', '附件')
RADIUS_PATH = os.path.join(OFFICIAL, '附件2.xlsx')
TEMPLATE_PATH = os.path.join(OFFICIAL, '附件3', 'result4.xlsx')
OUTPUT_DATA = os.path.join(PROJECT_ROOT, 'outputs', 'q4', 'data')
OUTPUT_WORKBOOK = os.path.join(PROJECT_ROOT, 'outputs', 'q4', 'workbooks',
                               'result4.xlsx')

C_TARGET = 0.15       # kg/kg  物理终止阈值，严格小于（题面"低于"）
C_DISP = 0.14995      # kg/kg  工作簿导出的保守判据，保证四位小数显示仍 < 0.1500
OUT_DT = 60.0         # s      result4 要求的输出间隔
R_OUT_CM = np.round(np.arange(21) * 0.1, 1)   # 0.0, 0.1, ..., 2.0 cm 固定物理距离列


class P4:
    """问题4 参数。几何与界面系数见 docs/q4 §4、§5 第11条。"""
    R0 = 0.02        # m          初始半径（附件2 的 t=0 值，也用于固定半径对照）
    L = 0.25         # m          圆柱长度，保持不变（一维径向模型中不参与计算）
    h = 25.0         # W/(m^2 K)  表面对流换热系数，沿用问题1
    hm = 8e-7        # m/s        表面传质系数，沿用问题1
    T0 = 28.0        # °C         初始温度
    C0 = 2.55        # kg/kg      初始干基含水率
    N = 200          # 材料坐标上的控制体数；Kirchhoff 界面下 N=200 已网格无关
    tol = P.tol      # Picard 收敛判据，与 q2/q3 一致
    maxit = P.maxit


# ---------------------------------------------------------------------------
# 物性：附录4 为正式口径，附录3 仅用于"同一离散下复算问题3"的对照
# ---------------------------------------------------------------------------
def rho4(C):
    return 760.0 + 90.0 * C


def cp4(C):
    return 1850.0 + 2150.0 * C / (C + 1.0)


def k4(C):
    return 0.12 + 0.20 * C / (C + 1.0)


# D(C,T_K) = D0 * exp(-aC/C) * exp(-Ea/T_K)，两套附录只差 D0 与 aC
PROPS = {
    'app4': dict(rho=rho4, cp=cp4, k=k4, D0=4.2e-4, aC=0.30, Ea=3850.0),
    'app3': dict(rho=rho_f, cp=cp_f, k=k_f, D0=2.4e-3, aC=0.45, Ea=3850.0),
}


def D_point(C, T_c, aC, D0, Ea):
    """点值扩散系数，只用于诊断与物性打印；界面通量一律走 Kirchhoff 平均。"""
    return (D0 * np.exp(-aC / np.maximum(C, 1e-12))
            * np.exp(-Ea / np.maximum(T_c + 273.15, 1.0)))


# ---------------------------------------------------------------------------
# Kirchhoff 积分平均界面扩散系数（PROJECT.md §12 第7条、docs/q4 §16.2）
# ---------------------------------------------------------------------------
def kirchhoff_psi(C, aC):
    """Kirchhoff 势 Psi(C) = \\int_0^C exp(-aC/s) ds 的解析式。

    推导：令 a = aC。由 d/ds [s e^{-a/s}] = e^{-a/s}(1 + a/s) 与
    \\int e^{-a/s}/s ds = E1(a/s)（换元 u=a/s），得
        Psi(C) = C e^{-a/C} - a E1(a/C),
    且 Psi(0+) = 0、Psi'(C) = e^{-a/C}，与被积函数一致。
    用解析式而非数值积分表，可以避免查找表的插值误差与额外标定参数。
    """
    Cs = np.maximum(np.asarray(C, dtype=float), 1e-12)
    x = aC / Cs
    return Cs * np.exp(-x) - aC * exp1(x)


def D_face_kirchhoff(C, T_c, aC, D0, Ea):
    """内部界面的 Kirchhoff 平均扩散系数，长度 N（节点数 N+1）。

    D_face = A(T_face) * [Psi(C_{i+1}) - Psi(C_i)] / (C_{i+1} - C_i)，
    对一维稳态通量精确：它等价于对非线性扩散势 A(T)Psi(C) 直接作差，因而能反映
    一个网格单元内跨越较大含水率区间的**平均**传输能力。点值调和平均会被单元内
    最小的 D 支配，在 D 随低含水率指数坍塌时给出近似为零的界面系数。
    A(T) 用界面两侧温度的算术平均：T 在本问中空间近乎均匀，取法不敏感。
    """
    A_face = D0 * np.exp(-Ea / (0.5 * (T_c[:-1] + T_c[1:]) + 273.15))
    psi = kirchhoff_psi(C, aC)
    dC = C[1:] - C[:-1]
    # |dC| 过小时差商退化为 0/0，取极限值 D(C_mid)；正常网格上不会触发
    tiny = np.abs(dC) < 1e-10
    ratio = np.where(
        tiny,
        np.exp(-aC / np.maximum(0.5 * (C[:-1] + C[1:]), 1e-12)),
        (psi[1:] - psi[:-1]) / np.where(tiny, 1.0, dC),
    )
    return A_face * ratio


# ---------------------------------------------------------------------------
# 附件2：半径历史
# ---------------------------------------------------------------------------
def load_radius(path=None):
    """读取并校验附件2，返回 (R_of_t[m], t_last[s], R_last[m])。

    附件2 的半径单位是 cm，进入方程前必须换成 m（docs/q4 §20.1）。
    t > t_last 时保持末值：主模型不外推收缩，只在结果中报告发生了外推。
    只用 R(t) 本身、不用 Rdot，因此节点处斜率不连续不影响正式方程。
    """
    if path is None:
        path = RADIUS_PATH
    d = pd.read_excel(path).values.astype(float)
    t, R_cm = d[:, 0], d[:, 1]
    if not np.isfinite(t).all() or not np.isfinite(R_cm).all():
        raise ValueError('附件2 存在缺失或非数值数据')
    if not np.all(np.diff(t) > 0):
        raise ValueError('附件2 时间非严格递增')
    if not np.all(np.diff(R_cm) <= 0):
        raise ValueError('附件2 半径出现回升，与单调收缩假设矛盾')
    R_m = R_cm / 100.0
    f = interp1d(t, R_m, 'linear', bounds_error=False,
                 fill_value=(R_m[0], R_m[-1]))
    return f, float(t[-1]), float(R_m[-1])


def constant_radius(R):
    """固定半径工况。用于 (a) 退化检验，(b) docs/q4 §19 的收缩效应对照。"""
    return lambda t: R


def dt_schedule(t, Cmax, scale=1.0):
    """变步长。所有取值除以 scale 后仍须整除 OUT_DT=60，保证输出时刻精确落在
    步点上，无需时间插值；scale 取 1/2/4 用于时间步收敛试验。

    阈值附近强制细化到 1 s：终止事件要定位到秒级，并且导出判据 0.14995 与物理
    判据 0.15 之间只差 5e-5，必须用小步长跨过。
    """
    if Cmax < 1.01 * C_TARGET:
        return 1.0 / scale       # 阈值附近：事件定位
    if t < 600.0:
        return 1.0 / scale       # 预热初期：边界阶跃、梯度最陡
    if t < 3600.0:
        return 5.0 / scale
    if t < 14400.0:
        return 20.0 / scale      # 预热结束，进入恒温段
    return 60.0 / scale          # 恒温干燥期：R^2/D 量级为数十小时


# ---------------------------------------------------------------------------
# 正式求解器
# ---------------------------------------------------------------------------
def solve_q4(R_of, props='app4', N=None, dt_scale=1.0, t_max=150 * 3600.0,
             env=None, h=None, hm=None, keep_full_dt=600.0, verbose=True):
    """在材料坐标 xi 上求解移动边界热-湿耦合问题。

    离散推导（对应 docs/q4 §16.1，这里是顶点型的版本）：
        固定域水分方程   dC/dt = 1/(R^2 xi) d/dxi [xi D dC/dxi]
        对控制体积分并两边乘 R(t)^2：
            R^2 W_i (C_i^{n+1} - C_i^n)/dt = G_{i+1/2} - G_{i-1/2},
            G_{i+1/2} = xi_{i+1/2} D_{i+1/2} (C_{i+1} - C_i)/dxi,
        其中 W_i = \\int xi dxi 为材料面积权重，sum(W) = 1/2。
        表面 (i=N) 用 Robin 条件 -D_s/R dC/dxi|_1 = hm (C_s - C_inf) 得
            G_{N+1/2} = -R hm (C_N - C_inf)。
        于是矩阵结构与固定半径问题**完全相同**：只要把质量项取成 R(t)^2 * W，
        表面系数取成 s = hm*R(t)，就能直接复用 q2_coupled.assemble/bsolve。
        热方程同理，质量项为 rho*cp*W*R^2，表面系数 s = h*R(t)。
        R 取新时间层的值 R(t^{n+1})，与后向欧拉一致。

    参数
        R_of        : 可调用对象，t[s] -> R[m]
        props       : 'app4'（正式）或 'app3'（同一离散下复算问题3 的对照）
        dt_scale    : 步长缩放，见 dt_schedule
        env         : (T_inf, C_inf) 插值函数对；None 时用问题3 的平台均值延拓
        h, hm       : None 时取题给值；置 0 用于绝热/零传质极限测试
        keep_full_dt: 完整 xi 剖面的保存间隔，供论文时空图使用；None 表示不保存
    返回 dict，键见函数末尾。
    """
    pr = PROPS[props]
    N = P4.N if N is None else N
    h = P4.h if h is None else h
    hm = P4.hm if hm is None else hm
    T_inf, C_inf = env if env is not None else load_env_ext()[:2]

    # 材料坐标网格：make_grid(N, 1.0) 给出 xi 节点、面积权重 W 和界面位置
    dxi, xi, W, xif = make_grid(N, 1.0)
    T = np.full(N + 1, P4.T0)
    C = np.full(N + 1, P4.C0)
    t = 0.0
    Cbar0 = 2.0 * (W * C).sum()          # 材料面积权重下的平均含水率，初值 = C0
    flux_cum = 0.0                        # 累计 \\int 2 hm/R (C_s - C_inf) dt

    r_out_m = R_OUT_CM / 100.0
    rec = dict(t=[], C=[], Cs=[], R=[], Cmax=[], Cbar=[], Tc=[])
    full = dict(t=[], C=[], T=[], R=[])
    t_f = t_disp = None
    itmax, m_res, e_res, checked = 0, 0.0, 0.0, False

    while t < t_max:
        dt = dt_schedule(t, C.max(), dt_scale)
        t1 = t + dt
        R1 = float(R_of(t1))
        To, Co = T.copy(), C.copy()
        Tn, Cn = T.copy(), C.copy()

        for it in range(P4.maxit):
            # (a) 传质：界面 D 用 Kirchhoff 平均，系数在时间中点取值
            Df = D_face_kirchhoff(0.5 * (Co + Cn), 0.5 * (To + Tn),
                                  pr['aC'], pr['D0'], pr['Ea'])
            lo, di, up = assemble(Df, xif, dxi, N, hm * R1)
            MC = W * R1 ** 2 / dt
            bC = MC * Co
            bC[-1] += hm * R1 * C_inf(t1)
            if not checked:
                check_K(lo, di, up, hm * R1, N)
            C2 = bsolve(lo, di, up, MC, bC, N)

            # (b) 传热：物性用刚更新的 C；k 界面取算术平均
            Cm = 0.5 * (Co + C2)
            MT = pr['rho'](Cm) * pr['cp'](Cm) * W * R1 ** 2 / dt
            lo2, di2, up2 = assemble(arith(pr['k'](Cm)), xif, dxi, N, h * R1)
            bT = MT * To
            bT[-1] += h * R1 * T_inf(t1)
            if not checked:
                check_K(lo2, di2, up2, h * R1, N)
                checked = True
            T2 = bsolve(lo2, di2, up2, MT, bT, N)

            err = max(np.abs(T2 - Tn).max(), np.abs(C2 - Cn).max())
            Tn, Cn = T2, C2
            if err < P4.tol:
                break
        else:
            warnings.warn(f'Picard 未收敛：t={t1:.1f}s，残差={err:.3e}，'
                          f'已达上限 maxit={P4.maxit}')
        itmax = max(itmax, it + 1)

        # 移动域下的守恒残差：dCbar/dt = -2 hm/R (C_s - C_inf)（docs/q4 §15.1）
        # 注意用材料面积权重 W，不能用随 R^2 缩小的几何体积，否则纯收缩会被误判为失水
        m_res = max(m_res, abs((W * (Cn - Co)).sum() / dt
                               + hm / R1 * (Cn[-1] - C_inf(t1))))
        # MT 已含 1/dt，这里不再除 dt
        e_res = max(e_res, abs((MT * (Tn - To)).sum()
                               - h * R1 * (T_inf(t1) - Tn[-1])))
        flux_cum += 2.0 * hm / R1 * (Cn[-1] - C_inf(t1)) * dt

        T, C, t = Tn, Cn, t1

        # --- 记录 60s 输出网格。t 由 1/5/20/60 等整除 60 的步长累加，取模精确 ---
        if abs(t % OUT_DT) < 1e-9:
            # 材料坐标 -> 物理位置：xi_q = r_q / R(t)，超出当前半径的位置记 NaN，
            # 既不外推也不填 0（docs/q4 §18.1、§18.2）
            xq = r_out_m / R1
            inside = xq <= 1.0 + 1e-12
            row = np.where(inside, np.interp(np.minimum(xq, 1.0), xi, C), np.nan)
            rec['t'].append(t)
            rec['C'].append(row)
            rec['Cs'].append(C[-1])          # 顶点型网格上表面就是节点，无需重构
            rec['R'].append(R1)
            rec['Cmax'].append(C.max())
            rec['Cbar'].append(2.0 * (W * C).sum())
            rec['Tc'].append(T[0])
        if keep_full_dt is not None and abs(t % keep_full_dt) < 1e-9:
            full['t'].append(t)
            full['C'].append(C.copy())
            full['T'].append(T.copy())
            full['R'].append(R1)

        # --- 终止事件。用全域 max 而非中心值：剖面一旦非单调可立即暴露 ---
        if t_f is None and C.max() < C_TARGET:
            t_f = t                          # 步长已细化到 1s，事件定位到秒级
        if t_f is not None and C.max() < C_DISP and abs(t % OUT_DT) < 1e-9:
            t_disp = t                       # 四位小数下仍严格 < 0.1500 的首个输出时刻
            break
        if verbose and abs(t % (6 * 3600.0)) < 1e-9:
            print(f'    t={t / 3600:6.1f}h  R={R1 * 100:.3f}cm  '
                  f'C_max={C.max():.5f}  C_表面={C[-1]:.5f}  T_中心={T[0]:.3f}')

    if t_disp is None and hm > 0.0:      # hm=0 的极限测试本就不会干完，不告警
        warnings.warn(f'到 t_max={t_max / 3600:.1f}h 仍未达到导出判据 '
                      f'C<{C_DISP}，当前 C_max={C.max():.5f}')

    return dict(
        t_f=t_f, t_disp=t_disp, t=np.array(rec['t']),
        C=np.array(rec['C']), Cs=np.array(rec['Cs']), R=np.array(rec['R']),
        Cmax=np.array(rec['Cmax']), Cbar=np.array(rec['Cbar']),
        Tc=np.array(rec['Tc']), r_cm=R_OUT_CM.copy(),
        full_t=np.array(full['t']), full_C=np.array(full['C']),
        full_T=np.array(full['T']), full_R=np.array(full['R']),
        xi=xi, W=W, C_fin=C, T_fin=T, N=N,
        Cbar0=Cbar0, Cbar_end=2.0 * (W * C).sum(), flux_cum=flux_cum,
        itmax=itmax, m_res=m_res, e_res=e_res)


# ---------------------------------------------------------------------------
# result4.xlsx 导出与回读验收（docs/q4 §18.2、§20.6）
# ---------------------------------------------------------------------------
def export_result4(res, template_path=TEMPLATE_PATH, out_path=OUTPUT_WORKBOOK):
    """在官方模板副本上写入完整含水率结果，随后回读验收。

    列结构：A=时间/s，B..V=到中心 0.0~2.0cm（单位 cm），W=药材表面。
    r > R(t) 的单元格留空，表示该位置已在药材外部。
    只在写入这一步四舍五入到四位小数；求解与事件定位全程使用未舍入值。
    """
    from openpyxl import load_workbook

    t_s, tab, Cs, r_cm, R_m = res['t'], res['C'], res['Cs'], res['r_cm'], res['R']
    nrow, ncol = 1 + len(t_s), 2 + len(r_cm)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    shutil.copyfile(template_path, out_path)          # 官方原始文件只读，改副本
    wb = load_workbook(out_path)
    ws = wb[wb.sheetnames[0]]

    for j, r in enumerate(r_cm):                      # 表头保留 A1 的官方文字
        ws.cell(1, 2 + j, float(r))
    ws.cell(1, ncol, '药材表面')
    for i in range(len(t_s)):
        ws.cell(2 + i, 1, float(t_s[i]))
        for j in range(len(r_cm)):
            v = tab[i, j]
            ws.cell(2 + i, 2 + j, None if not np.isfinite(v) else round(float(v), 4))
        ws.cell(2 + i, ncol, round(float(Cs[i]), 4))
    if ws.max_row > nrow:                             # 清掉模板里的 "…" 占位行
        ws.delete_rows(nrow + 1, ws.max_row - nrow)
    wb.save(out_path)
    wb.close()
    return validate_result4(out_path, t_s, r_cm, R_m)


def validate_result4(path, t_s, r_cm, R_m):
    """回读工作簿，检查尺寸、表头、时间列、留空区域与末行显示值。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    nrow, ncol = 1 + len(t_s), 2 + len(r_cm)
    assert ws.max_row == nrow, f'行数应为 {nrow}，实际 {ws.max_row}'
    assert ws.max_column == ncol, f'列数应为 {ncol}，实际 {ws.max_column}'
    assert ws.cell(1, ncol).value == '药材表面', '最后一列表头应为"药材表面"'

    n_num = n_blank = 0
    for i in range(len(t_s)):
        assert float(ws.cell(2 + i, 1).value) == float(t_s[i]), f'第{i + 2}行时间不符'
        inside = r_cm / 100.0 <= R_m[i] * (1.0 + 1e-12)
        for j in range(len(r_cm)):
            v = ws.cell(2 + i, 2 + j).value
            if inside[j]:
                assert isinstance(v, (int, float)), \
                    f't={t_s[i]}s, r={r_cm[j]}cm 应为数值，实际 {v!r}'
                n_num += 1
            else:
                assert v is None, \
                    f't={t_s[i]}s, r={r_cm[j]}cm 已在药材外部，应留空，实际 {v!r}'
                n_blank += 1
        assert isinstance(ws.cell(2 + i, ncol).value, (int, float)), '表面列应为数值'
    last = [ws.cell(nrow, 2 + j).value for j in range(len(r_cm) + 1)]
    worst = max(v for v in last if v is not None)
    assert worst < C_TARGET, f'末行显示值 {worst} 未严格低于 0.1500'
    wb.close()
    return dict(rows=nrow, cols=ncol, numeric=n_num, blank=n_blank, last_max=worst)


# ---------------------------------------------------------------------------
def _hours(x):
    return '—' if x is None else f'{x / 3600:.3f} h'


if __name__ == '__main__':
    quick = '--quick' in sys.argv
    print('=' * 74)
    print('问题4：附件2 收缩条件下的烘干时长（材料坐标 + Kirchhoff 界面）')
    print('=' * 74)

    R_of, t_last, R_last = load_radius()
    T_inf, C_inf, Te, Ce = load_env_ext()
    print(f'附件2：0~{t_last / 3600:.0f}h，R: {R_of(0.0) * 100:.3f} -> '
          f'{R_last * 100:.3f} cm，体积收缩至 {(R_last / P4.R0) ** 2 * 100:.2f}%')
    print(f'环境延拓：T∞={Te:.4f} °C，C∞={Ce:.5f} kg/kg（t≥9000s 平台均值）')
    print(f'收缩带来的内部传输加速因子 (R0/R_last)^2 = '
          f'{(P4.R0 / R_last) ** 2:.2f}')

    if not quick:
        print('\n--- 极限测试（docs/q4 §20.4）---')
        r0 = solve_q4(R_of, hm=0.0, t_max=6 * 3600.0, keep_full_dt=None,
                      verbose=False)
        print(f'  hm=0：半径 2.000 -> {r0["R"][-1] * 100:.3f} cm，'
              f'平均含水率漂移 {abs(r0["Cbar_end"] - r0["Cbar0"]):.3e} kg/kg')
        assert abs(r0['Cbar_end'] - r0['Cbar0']) < 1e-12, \
            '收缩被误算成失水：材料面积权重下 Cbar 必须守恒'
        r1 = solve_q4(R_of, h=0.0, hm=0.0, t_max=3600.0, keep_full_dt=None,
                      verbose=False)
        assert np.allclose(r1['T_fin'], P4.T0) and np.allclose(r1['C_fin'], P4.C0), \
            'h=hm=0 时均匀初态必须保持不变'
        print('  h=hm=0：温度与含水率场保持均匀初态，通过')

        print('\n--- 网格与时间步收敛（烘干时间 / h）---')
        conv = {}
        for Nx in [100, 200, 400]:
            rr = solve_q4(R_of, N=Nx, keep_full_dt=None, verbose=False)
            conv[Nx] = rr['t_f']
            print(f'  N={Nx:4d}  dt×1   {_hours(conv[Nx])}')
        rr = solve_q4(R_of, dt_scale=2.0, keep_full_dt=None, verbose=False)
        print(f'  N={P4.N:4d}  dt×1/2 {_hours(rr["t_f"])}')

    print('\n--- 正式求解（附录4 物性 + 附件2 半径）---')
    res = solve_q4(R_of)
    te, td = res['t_f'], res['t_disp']
    i60 = int(np.argmax(res['Cmax'] < C_TARGET))
    t_f60 = res['t'][i60]
    print(f'\n  连续事件时间 t_f      = {te:.0f} s = {_hours(te)} = {te / 86400:.3f} 天')
    print(f'  60s 网格首次达标      = {t_f60:.0f} s = {_hours(t_f60)}')
    print(f'  工作簿保守判据 (<{C_DISP}) = {td:.0f} s = {_hours(td)}')
    if td > t_last:
        print(f'  * 注意：终止时刻超过附件2 的 {t_last / 3600:.0f}h，'
              f'该段半径按末值 {R_last * 100:.3f} cm 外推')
    print(f'  Picard 最大迭代 = {res["itmax"]}，'
          f'质量守恒残差 = {res["m_res"]:.2e}，能量守恒残差 = {res["e_res"]:.2e}')
    print(f'  累计水分守恒 Cbar(t)-Cbar(0)+∫2hm/R(Cs-C∞)dt = '
          f'{res["Cbar_end"] - res["Cbar0"] + res["flux_cum"]:.3e}')
    print(f'  终态 C ∈ [{res["C_fin"].min():.5f}, {res["C_fin"].max():.5f}]，'
          f'T ∈ [{res["T_fin"].min():.4f}, {res["T_fin"].max():.4f}] °C')

    # --- 物理自检（docs/q4 §20.3）---
    C = res['C']
    assert np.all(np.diff(res['Cmax']) < 1e-9), '全域最大含水率非单调下降'
    assert np.all(res['Cs'] <= C[:, 0] + 1e-9), '表面应始终不湿于中心'
    assert np.nanmin(C) > 0.0, '含水率出现非正值'
    assert np.all(np.isnan(C) == (res['r_cm'][None, :] / 100.0
                                  > res['R'][:, None] * (1 + 1e-12))), \
        '留空区域与当前半径不一致'
    assert res['Cmax'][-1] < C_DISP and res['Cmax'][i60 - 1] >= C_TARGET, \
        '终止时刻不是首次达标'
    assert np.all(np.diff(res['R']) <= 1e-15), '半径出现回升'
    print('  物理自检通过')

    print('\n--- 收缩效应分解（docs/q4 §19，三组模型同一 Kirchhoff 离散）---')
    ctrl4 = solve_q4(constant_radius(P4.R0), keep_full_dt=None, verbose=False)
    ctrl3 = solve_q4(constant_radius(P4.R0), props='app3', keep_full_dt=None,
                     verbose=False)
    print(f'  附录3 物性 + R=2cm（问题3 口径）      {_hours(ctrl3["t_f"])}')
    print(f'  附录4 物性 + R=2cm（固定半径对照）    {_hours(ctrl4["t_f"])}')
    print(f'  附录4 物性 + 附件2 R(t)（问题4 正式）  {_hours(te)}')
    print(f'  物性差异贡献 {(ctrl4["t_f"] - ctrl3["t_f"]) / 3600:+.3f} h，'
          f'收缩贡献 {(te - ctrl4["t_f"]) / 3600:+.3f} h')
    print('  注：附录3 行用 Kirchhoff 界面重算，与 q3_drying_time.py 的调和平均'
          '结果不可直接比较。')

    if not quick:
        print('\n--- 边界延拓敏感性（附件末值代替平台均值，仅作对照）---')
        env_last = load_env_ext(plateau_t=14400.0)   # 只选中末点，等价于末值延拓
        sens = solve_q4(R_of, env=env_last[:2], keep_full_dt=None, verbose=False)
        print(f'  平台均值延拓 {_hours(te)}   末值延拓 {_hours(sens["t_f"])}   '
              f'相对差 {abs(sens["t_f"] - te) / te * 100:.2f}%')

    # --- 结构化结果：未舍入，供论文图与后续复算使用（不从 Excel 反向作图）---
    # 命名必须在**忽略大小写**后仍唯一：macOS 的 APFS 默认大小写不敏感，
    # Q4_R.npy / Q4_r.npy 会指向同一个文件，后写的直接覆盖先写的。
    dumps = {
        'Q4_time.npy': res['t'],          # 60s 输出时刻 / s
        'Q4_C.npy': res['C'],             # 固定物理距离上的含水率，体外为 NaN
        'Q4_Csurf.npy': res['Cs'],        # 药材表面含水率
        'Q4_Rhist.npy': res['R'],         # 各输出时刻的半径 / m
        'Q4_rout.npy': res['r_cm'],       # 固定物理距离列 / cm
        'Q4_xi.npy': res['xi'],           # 材料坐标节点
        'Q4_full_time.npy': res['full_t'],
        'Q4_full_C.npy': res['full_C'],   # 完整 xi 剖面（含水率）
        'Q4_full_Temp.npy': res['full_T'],
        'Q4_full_Rhist.npy': res['full_R'],
    }
    lowered = [n.lower() for n in dumps]
    assert len(set(lowered)) == len(lowered), '输出文件名在忽略大小写后不唯一'
    os.makedirs(OUTPUT_DATA, exist_ok=True)
    for name, arr in dumps.items():
        np.save(os.path.join(OUTPUT_DATA, name), arr)
    print(f'\n未舍入结果已保存到 {OUTPUT_DATA}')

    # --- 表6：每 6h + 结束时刻，每 0.5cm + 药材表面 ---
    # 末行取物理首次达标时刻 t_f60，并按 6 位小数打印原始值：中心的 0.149977 若按
    # 4 位小数显示会变成 0.1500，看上去违反"低于 0.15"（result3.xlsx 当初就栽在
    # 这里）。论文表不受题目 4 位小数的约束，直接如实给出真值即可消除歧义；
    # result4.xlsx 仍按 4 位小数、并写到更保守的 t_disp（PROJECT.md §11.4）。
    cols = [0, 5, 10, 15, 20]
    print('\n表6  收缩条件下药材的水分浓度 (kg/kg)')
    print('  时间/h |' + ''.join(f'{res["r_cm"][c]:>11.1f}' for c in cols)
          + '      药材表面      R/cm')
    marks = list(np.arange(6.0, t_f60 / 3600.0, 6.0)) + [t_f60 / 3600.0]
    for th in marks:
        k = min(int(round(th * 3600.0 / OUT_DT)) - 1, len(res['t']) - 1)
        cells = ''.join('          —' if not np.isfinite(res['C'][k, c])
                        else f'{res["C"][k, c]:11.6f}' for c in cols)
        print(f'{th:8.3f} |' + cells
              + f'{res["Cs"][k]:14.6f}{res["R"][k] * 100:10.3f}')
    assert max(res['C'][i60, c] for c in cols
               if np.isfinite(res['C'][i60, c])) < C_TARGET, \
        '表6 末行原始值未严格低于 0.15'
    print('  "—" 表示该固定位置此刻已在药材外部。')
    print(f'  末行 {t_f60 / 3600:.3f} h 为 60s 输出网格上首次全域低于 0.15 的时刻，'
          f'按原始值给出；连续模型事件时间 {te / 3600:.3f} h。')
    print(f'  result4.xlsx 按题目要求保留 4 位小数，因此写到更保守的 '
          f'{td / 3600:.3f} h（原始值 < {C_DISP}）。')

    # --- 官方模板导出 ---
    chk = export_result4(res)
    print(f'\nresult4.xlsx 已写入 {OUTPUT_WORKBOOK}')
    print(f'  回读验收：{chk["rows"]} 行 × {chk["cols"]} 列，'
          f'有效数值 {chk["numeric"]} 个，表面外留空 {chk["blank"]} 个，'
          f'末行最大显示值 {chk["last_max"]:.4f}')

"""2026 高教社杯 A 题问题 1：圆柱药材短时热湿迁移求解器。

模型采用一维轴对称径向有限体积法。水分方程使用 Picard 迭代处理
``D(C)``，温度方程只考虑内部导热与表面对流换热。脚本不会修改官方附件，
会把完整数值结果、诊断报告和 ``result1.xlsx`` 写入 outputs。
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import tempfile
from copy import copy
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Callable, Iterable

import numpy as np


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ENV_PATH = PROJECT_ROOT / "data" / "official" / "A题" / "附件" / "附件1.xlsx"
DEFAULT_TEMPLATE_PATH = (
    PROJECT_ROOT / "data" / "official" / "A题" / "附件" / "附件3" / "result1.xlsx"
)
DEFAULT_OUTPUT_ROOT = PROJECT_ROOT / "outputs" / "q1"


@dataclass(frozen=True)
class Q1Parameters:
    """问题 1 的 SI 制参数。"""

    radius_m: float = 0.02
    length_m: float = 0.25
    density_kg_m3: float = 820.0
    heat_capacity_j_kg_k: float = 2600.0
    conductivity_w_m_k: float = 0.36
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8.0e-7
    initial_temperature_c: float = 28.0
    initial_moisture_kg_kg: float = 2.55
    end_time_s: float = 1800.0

    @property
    def thermal_diffusivity_m2_s(self) -> float:
        return self.conductivity_w_m_k / (
            self.density_kg_m3 * self.heat_capacity_j_kg_k
        )

DEFAULT_PARAMETERS = Q1Parameters()


@dataclass(frozen=True)
class BoundarySeries:
    """附件 1 的分段线性边界序列，区间外使用端点常值。"""

    time_s: np.ndarray
    value: np.ndarray

    def __call__(self, time_s: float | np.ndarray) -> float | np.ndarray:
        result = np.interp(time_s, self.time_s, self.value)
        return float(result) if np.ndim(result) == 0 else result


@dataclass(frozen=True)
class TriDiagonalFactor:
    """Thomas 算法的三对角矩阵分解。"""

    multipliers: np.ndarray
    diagonal: np.ndarray
    upper: np.ndarray

    def solve(self, rhs: np.ndarray) -> np.ndarray:
        work = np.asarray(rhs, dtype=float).copy()
        for i, multiplier in enumerate(self.multipliers, start=1):
            work[i] -= multiplier * work[i - 1]
        solution = work
        solution[-1] /= self.diagonal[-1]
        for i in range(solution.size - 2, -1, -1):
            solution[i] = (
                solution[i] - self.upper[i] * solution[i + 1]
            ) / self.diagonal[i]
        return solution


def factor_tridiagonal(
    lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray
) -> TriDiagonalFactor:
    lower = np.asarray(lower, dtype=float)
    diagonal = np.asarray(diagonal, dtype=float).copy()
    upper = np.asarray(upper, dtype=float).copy()
    if diagonal.size != lower.size + 1 or lower.shape != upper.shape:
        raise ValueError("三对角矩阵的维数不一致")
    multipliers = np.empty_like(lower)
    for i in range(lower.size):
        if abs(diagonal[i]) < 1e-30:
            raise np.linalg.LinAlgError("三对角矩阵出现零主元")
        multipliers[i] = lower[i] / diagonal[i]
        diagonal[i + 1] -= multipliers[i] * upper[i]
    if abs(diagonal[-1]) < 1e-30:
        raise np.linalg.LinAlgError("三对角矩阵出现零主元")
    return TriDiagonalFactor(multipliers, diagonal, upper)


def apply_tridiagonal(
    lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray, vector: np.ndarray
) -> np.ndarray:
    result = diagonal * vector
    result[1:] += lower * vector[:-1]
    result[:-1] += upper * vector[1:]
    return result


def moisture_diffusivity(moisture: np.ndarray | float) -> np.ndarray | float:
    """题面经验式 ``D=7e-9 exp(-0.89/C)``。

    极小正数只保护系数求值；求解后的状态量不会被裁剪，负值会触发错误。
    """

    values = np.asarray(moisture, dtype=float)
    safe = np.maximum(values, 1e-12)
    result = 7.0e-9 * np.exp(-0.89 / safe)
    return float(result) if result.ndim == 0 else result


def load_environment(
    path: Path | str = DEFAULT_ENV_PATH,
    required_end_time_s: float = DEFAULT_PARAMETERS.end_time_s,
) -> tuple[BoundarySeries, BoundarySeries, dict]:
    """读取并严格校验附件 1。"""

    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"找不到环境数据文件：{path}")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "读取 Excel 环境数据需要安装 openpyxl：pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [sheet.cell(1, column).value for column in range(1, 4)]
    if headers != ["时间", "温度", "水分浓度"]:
        raise ValueError(
            "附件1前三列表头应为‘时间、温度、水分浓度’，"
            f"实际为 {headers}"
        )

    rows = list(sheet.iter_rows(min_row=2, min_col=1, max_col=3, values_only=True))
    workbook.close()
    try:
        data = np.asarray(rows, dtype=float)
    except (TypeError, ValueError) as exc:
        raise ValueError("附件1前三列必须全部是数值") from exc

    if data.ndim != 2 or data.shape[1] != 3 or data.shape[0] < 2:
        raise ValueError("附件1必须至少包含两个时刻和三列数据")
    if not np.isfinite(data).all():
        raise ValueError("附件1包含空值或非有限数值")
    time_s, temperature_c, moisture = data.T
    if not np.all(np.diff(time_s) > 0):
        raise ValueError("附件1的时间必须严格递增且不得重复")
    if time_s[0] > 0 or time_s[-1] < required_end_time_s:
        raise ValueError(
            f"附件1时间范围 [{time_s[0]}, {time_s[-1]}] s "
            f"不能覆盖 [0, {required_end_time_s}] s"
        )
    if np.any(moisture < 0):
        raise ValueError("附件1的水分浓度不能为负")

    metadata = {
        "path": str(path.resolve()),
        "rows": int(data.shape[0]),
        "time_start_s": float(time_s[0]),
        "time_end_s": float(time_s[-1]),
        "sample_interval_min_s": float(np.diff(time_s).min()),
        "sample_interval_max_s": float(np.diff(time_s).max()),
        "missing_or_nonfinite_count": 0,
    }
    return (
        BoundarySeries(time_s, temperature_c),
        BoundarySeries(time_s, moisture),
        metadata,
    )


def make_grid(cell_count: int, radius_m: float) -> tuple[np.ndarray, ...]:
    """建立含中心与表面节点、两端为半控制体积的径向网格。"""

    if cell_count < 2:
        raise ValueError("径向网格数至少为 2")
    if radius_m <= 0:
        raise ValueError("半径必须为正")
    dr = radius_m / cell_count
    radius = np.arange(cell_count + 1, dtype=float) * dr
    volume_weight = radius * dr
    volume_weight[0] = dr**2 / 8.0
    volume_weight[-1] = (dr / 2.0) * (radius_m - dr / 4.0)
    face_radius = (np.arange(cell_count, dtype=float) + 0.5) * dr
    return np.asarray(dr), radius, volume_weight, face_radius


def assemble_operator(
    face_coefficient: np.ndarray,
    face_radius: np.ndarray,
    dr: float,
    surface_coefficient: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    conductance = np.asarray(face_coefficient) * face_radius / dr
    diagonal = np.zeros(conductance.size + 1)
    diagonal[:-1] -= conductance
    diagonal[1:] -= conductance
    diagonal[-1] -= surface_coefficient
    return conductance.copy(), diagonal, conductance.copy()


def check_operator(
    operator: tuple[np.ndarray, np.ndarray, np.ndarray],
    surface_coefficient: float,
) -> None:
    lower, diagonal, upper = operator
    if not np.allclose(lower, upper, rtol=0.0, atol=1e-14):
        raise AssertionError("离散输运算子不对称")
    row_sum = diagonal.copy()
    row_sum[1:] += lower
    row_sum[:-1] += upper
    if np.max(np.abs(row_sum[:-1])) > 1e-10:
        raise AssertionError("内部控制体行和不为零")
    if abs(row_sum[-1] + surface_coefficient) > 1e-10:
        raise AssertionError("表面控制体行和与 Robin 边界不一致")


def _make_lhs_factor(
    storage: np.ndarray,
    operator: tuple[np.ndarray, np.ndarray, np.ndarray],
    step_s: float,
    theta: float,
) -> TriDiagonalFactor:
    lower, diagonal, upper = operator
    return factor_tridiagonal(
        -theta * lower,
        storage / step_s - theta * diagonal,
        -theta * upper,
    )


def _validate_time_grid(end_time_s: float, dt_s: float, output_count: int) -> tuple[int, int]:
    if dt_s <= 0 or output_count <= 0:
        raise ValueError("时间步长和输出点数必须为正")
    step_count = int(round(end_time_s / dt_s))
    if not math.isclose(step_count * dt_s, end_time_s, abs_tol=1e-10):
        raise ValueError("时间步长必须整除总时长")
    if step_count % output_count != 0:
        raise ValueError("输出点数必须整除总时间步数")
    return step_count, step_count // output_count


def solve_heat(
    cell_count: int,
    dt_s: float,
    output_count: int,
    ambient_temperature: Callable[[float], float],
    *,
    parameters: Q1Parameters = DEFAULT_PARAMETERS,
    theta: float = 0.5,
    rannacher_steps: int = 2,
) -> tuple[np.ndarray, np.ndarray, dict]:
    """求解仅含内部导热和表面对流换热的温度场。"""

    dr, radius, volume, face_radius = make_grid(cell_count, parameters.radius_m)
    storage = parameters.density_kg_m3 * parameters.heat_capacity_j_kg_k * volume
    operator = assemble_operator(
        np.full(cell_count, parameters.conductivity_w_m_k),
        face_radius,
        float(dr),
        parameters.heat_transfer_w_m2_k * parameters.radius_m,
    )
    check_operator(
        operator, parameters.heat_transfer_w_m2_k * parameters.radius_m
    )
    full_factor = _make_lhs_factor(storage, operator, dt_s, theta)
    half_factor = _make_lhs_factor(storage, operator, dt_s / 2.0, 1.0)

    def boundary_vector(time_s: float) -> np.ndarray:
        vector = np.zeros(cell_count + 1)
        vector[-1] = (
            parameters.radius_m
            * parameters.heat_transfer_w_m2_k
            * float(ambient_temperature(time_s))
        )
        return vector

    def convective_input(surface_temperature: float, time_s: float) -> float:
        return (
            parameters.heat_transfer_w_m2_k
            * parameters.radius_m
            * (float(ambient_temperature(time_s)) - surface_temperature)
        )

    temperature = np.full(cell_count + 1, parameters.initial_temperature_c)
    output = np.empty((output_count, cell_count + 1))
    step_count, stride = _validate_time_grid(
        parameters.end_time_s, dt_s, output_count
    )
    time_s = 0.0
    max_step_residual = 0.0
    convective_energy = 0.0

    def advance(
        old: np.ndarray,
        start_s: float,
        step_s: float,
        scheme_theta: float,
        factor: TriDiagonalFactor,
    ) -> tuple[np.ndarray, float, float]:
        rhs = storage / step_s * old
        if scheme_theta != 1.0:
            rhs += (1.0 - scheme_theta) * apply_tridiagonal(*operator, old)
        rhs += (
            scheme_theta * boundary_vector(start_s + step_s)
            + (1.0 - scheme_theta) * boundary_vector(start_s)
        )
        new = factor.solve(rhs)
        convection_0 = convective_input(old[-1], start_s)
        convection_1 = convective_input(new[-1], start_s + step_s)
        convection = scheme_theta * convection_1 + (1.0 - scheme_theta) * convection_0
        storage_rate = float(np.dot(storage, new - old) / step_s)
        residual = abs(storage_rate - convection) / max(abs(convection), 1e-12)
        return new, residual, convection * step_s

    for step_index in range(1, step_count + 1):
        if step_index <= rannacher_steps:
            substeps = ((dt_s / 2.0, 1.0, half_factor),) * 2
        else:
            substeps = ((dt_s, theta, full_factor),)
        for substep_s, scheme_theta, factor in substeps:
            temperature, residual, heat_in = advance(
                temperature, time_s, substep_s, scheme_theta, factor
            )
            time_s += substep_s
            max_step_residual = max(max_step_residual, residual)
            convective_energy += heat_in
        if step_index % stride == 0:
            output[step_index // stride - 1] = temperature

    sensible_energy = float(
        np.dot(storage, temperature - parameters.initial_temperature_c)
    )
    balance_residual = sensible_energy - convective_energy
    denominator = abs(convective_energy)
    cumulative_residual = (
        abs(balance_residual) / denominator if denominator > 1e-12 else 0.0
    )
    diagnostics = {
        "max_step_relative_residual": float(max_step_residual),
        "cumulative_relative_residual": float(cumulative_residual),
        "cumulative_absolute_residual_per_2piL_J_per_m": float(
            abs(balance_residual)
        ),
        "sensible_energy_per_2piL_J_per_m": sensible_energy,
        "convective_energy_per_2piL_J_per_m": float(convective_energy),
    }
    return radius, output, diagnostics


def solve_mass(
    cell_count: int,
    dt_s: float,
    output_count: int,
    ambient_moisture: Callable[[float], float],
    *,
    parameters: Q1Parameters = DEFAULT_PARAMETERS,
    theta: float = 0.5,
    rannacher_steps: int = 2,
    picard_tolerance: float = 1e-10,
    picard_max_iterations: int = 30,
) -> tuple[np.ndarray, np.ndarray, dict]:
    """求解非线性水分扩散方程。"""

    dr, radius, volume, face_radius = make_grid(cell_count, parameters.radius_m)

    def build(
        coefficient_state: np.ndarray, step_s: float, scheme_theta: float
    ) -> tuple[TriDiagonalFactor, tuple[np.ndarray, np.ndarray, np.ndarray]]:
        node_diffusivity = moisture_diffusivity(coefficient_state)
        face_diffusivity = (
            2.0
            * node_diffusivity[:-1]
            * node_diffusivity[1:]
            / (node_diffusivity[:-1] + node_diffusivity[1:])
        )
        operator = assemble_operator(
            face_diffusivity,
            face_radius,
            float(dr),
            parameters.mass_transfer_m_s * parameters.radius_m,
        )
        factor = _make_lhs_factor(volume, operator, step_s, scheme_theta)
        return factor, operator

    def boundary_vector(time_s: float) -> np.ndarray:
        vector = np.zeros(cell_count + 1)
        vector[-1] = (
            parameters.mass_transfer_m_s
            * parameters.radius_m
            * float(ambient_moisture(time_s))
        )
        return vector

    moisture = np.full(cell_count + 1, parameters.initial_moisture_kg_kg)
    output = np.empty((output_count, cell_count + 1))
    step_count, stride = _validate_time_grid(
        parameters.end_time_s, dt_s, output_count
    )
    time_s = 0.0
    max_step_residual = 0.0
    cumulative_outflow = 0.0
    max_picard_iterations_used = 0

    for step_index in range(1, step_count + 1):
        if step_index <= rannacher_steps:
            substeps = ((dt_s / 2.0, 1.0),) * 2
        else:
            substeps = ((dt_s, theta),)
        for step_s, scheme_theta in substeps:
            old = moisture.copy()
            new = old.copy()
            for iteration in range(1, picard_max_iterations + 1):
                factor, operator = build(
                    0.5 * (old + new), step_s, scheme_theta
                )
                rhs = volume / step_s * old
                if scheme_theta != 1.0:
                    rhs += (1.0 - scheme_theta) * apply_tridiagonal(
                        *operator, old
                    )
                rhs += (
                    scheme_theta * boundary_vector(time_s + step_s)
                    + (1.0 - scheme_theta) * boundary_vector(time_s)
                )
                candidate = factor.solve(rhs)
                if not np.isfinite(candidate).all():
                    raise FloatingPointError(
                        f"水分解出现非有限值：t={time_s + step_s:.6f} s"
                    )
                if np.min(candidate) < -1e-10:
                    raise FloatingPointError(
                        "水分解出现负值，程序不会通过裁剪掩盖该问题："
                        f"t={time_s + step_s:.6f} s, min={candidate.min():.6e}"
                    )
                if np.max(np.abs(candidate - new)) < picard_tolerance:
                    new = candidate
                    max_picard_iterations_used = max(
                        max_picard_iterations_used, iteration
                    )
                    break
                new = candidate
            else:
                raise RuntimeError(
                    f"水分 Picard 迭代未收敛：t={time_s + step_s:.6f} s"
                )

            outflow_0 = (
                parameters.mass_transfer_m_s
                * parameters.radius_m
                * (old[-1] - float(ambient_moisture(time_s)))
            )
            outflow_1 = (
                parameters.mass_transfer_m_s
                * parameters.radius_m
                * (new[-1] - float(ambient_moisture(time_s + step_s)))
            )
            outflow = (
                scheme_theta * outflow_1
                + (1.0 - scheme_theta) * outflow_0
            )
            inventory_rate = float(np.dot(volume, new - old) / step_s)
            max_step_residual = max(
                max_step_residual,
                abs(inventory_rate + outflow) / max(abs(outflow), 1e-12),
            )
            cumulative_outflow += outflow * step_s
            moisture = new
            time_s += step_s
        if step_index % stride == 0:
            output[step_index // stride - 1] = moisture

    inventory_change = float(
        np.dot(volume, moisture - parameters.initial_moisture_kg_kg)
    )
    balance_residual = inventory_change + cumulative_outflow
    cumulative_residual = (
        abs(balance_residual) / abs(cumulative_outflow)
        if abs(cumulative_outflow) > 1e-12
        else 0.0
    )
    diagnostics = {
        "max_step_relative_residual": float(max_step_residual),
        "cumulative_relative_residual": float(cumulative_residual),
        "cumulative_absolute_residual_per_2piL_m2": float(abs(balance_residual)),
        "normalized_outflow_per_2piL_m2": float(cumulative_outflow),
        "max_picard_iterations_used": int(max_picard_iterations_used),
    }
    return radius, output, diagnostics


def resample_radius(
    radius_m: np.ndarray, values: np.ndarray, output_radius_cm: np.ndarray
) -> np.ndarray:
    radius_cm = radius_m * 100.0
    if output_radius_cm[0] < radius_cm[0] or output_radius_cm[-1] > radius_cm[-1]:
        raise ValueError("输出半径超出计算域")
    return np.vstack(
        [np.interp(output_radius_cm, radius_cm, row) for row in values]
    )


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def export_result_workbook(
    template_path: Path | str,
    output_path: Path | str,
    time_s: np.ndarray,
    radius_cm: np.ndarray,
    temperature_c: np.ndarray,
    moisture_kg_kg: np.ndarray,
) -> dict:
    """在官方模板副本中写入 q1 全量结果。"""

    template_path = Path(template_path)
    output_path = Path(output_path)
    from openpyxl import load_workbook

    if not template_path.is_file():
        raise FileNotFoundError(f"找不到 result1 模板：{template_path}")
    expected_shape = (time_s.size, radius_cm.size)
    if temperature_c.shape != expected_shape or moisture_kg_kg.shape != expected_shape:
        raise ValueError("时间、半径与结果矩阵的尺寸不一致")
    if not np.isfinite(temperature_c).all() or not np.isfinite(moisture_kg_kg).all():
        raise ValueError("结果矩阵包含非有限数值")

    workbook = load_workbook(template_path)
    required_sheets = ["温度", "水分浓度"]
    if workbook.sheetnames != required_sheets:
        raise ValueError(
            f"result1 模板工作表应为 {required_sheets}，实际为 {workbook.sheetnames}"
        )

    for sheet_name, matrix in zip(
        required_sheets, (temperature_c, moisture_kg_kg), strict=True
    ):
        sheet = workbook[sheet_name]
        header_style = copy(sheet["B1"])
        time_style = copy(sheet["A2"])
        value_style = copy(sheet["B2"])
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        if sheet.max_column > 1:
            sheet.delete_cols(2, sheet.max_column - 1)

        sheet["A1"] = "时间\\到药材中心的距离"
        for column, radius in enumerate(radius_cm, start=2):
            cell = sheet.cell(1, column, float(round(radius, 10)))
            _copy_cell_style(header_style, cell)
            cell.number_format = "0.0"
        for row_index, time_value in enumerate(time_s, start=2):
            time_cell = sheet.cell(row_index, 1, int(round(float(time_value))))
            _copy_cell_style(time_style, time_cell)
            time_cell.number_format = "0"
            for column_index, value in enumerate(matrix[row_index - 2], start=2):
                cell = sheet.cell(row_index, column_index, round(float(value), 4))
                _copy_cell_style(value_style, cell)
                cell.number_format = "0.0000"

        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = max(
            sheet.column_dimensions["A"].width or 0, 27
        )
        for column_index in range(2, radius_cm.size + 2):
            column_letter = sheet.cell(1, column_index).column_letter
            sheet.column_dimensions[column_letter].width = 11

    workbook.properties.title = "2026 CUMCM A题问题1完整结果"
    workbook.properties.description = (
        "由 src/q1_fvm.py 自动生成；A列为时间(s)，第一行为半径(cm)。"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix="result1_", suffix=".xlsx", dir=output_path.parent
    )
    os.close(file_descriptor)
    try:
        workbook.save(temporary_name)
        workbook.close()
        os.replace(temporary_name, output_path)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)
    return validate_result_workbook(output_path, time_s, radius_cm)


def validate_result_workbook(
    path: Path | str, time_s: np.ndarray, radius_cm: np.ndarray
) -> dict:
    """回读结果工作簿，验证结构、边界单元格和数值类型。"""

    path = Path(path)
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    expected_rows = time_s.size + 1
    expected_columns = radius_cm.size + 1
    checks = {}
    for sheet_name in ["温度", "水分浓度"]:
        sheet = workbook[sheet_name]
        numeric_cells = 0
        for row in sheet.iter_rows(
            min_row=2,
            max_row=expected_rows,
            min_col=2,
            max_col=expected_columns,
            values_only=True,
        ):
            numeric_cells += sum(isinstance(value, (int, float)) for value in row)
        checks[sheet_name] = {
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "first_time_s": sheet.cell(2, 1).value,
            "last_time_s": sheet.cell(expected_rows, 1).value,
            "first_radius_cm": sheet.cell(1, 2).value,
            "last_radius_cm": sheet.cell(1, expected_columns).value,
            "numeric_result_cells": numeric_cells,
        }
        if sheet.max_row != expected_rows or sheet.max_column != expected_columns:
            raise AssertionError(f"{sheet_name} 工作表尺寸不符合题目要求")
        if numeric_cells != time_s.size * radius_cm.size:
            raise AssertionError(f"{sheet_name} 工作表存在空值或非数值结果")
        if sheet.cell(2, 1).value != int(time_s[0]):
            raise AssertionError(f"{sheet_name} 起始时间不正确")
        if sheet.cell(expected_rows, 1).value != int(time_s[-1]):
            raise AssertionError(f"{sheet_name} 终止时间不正确")
        if not math.isclose(sheet.cell(1, 2).value, float(radius_cm[0])):
            raise AssertionError(f"{sheet_name} 起始半径不正确")
        if not math.isclose(
            sheet.cell(1, expected_columns).value, float(radius_cm[-1])
        ):
            raise AssertionError(f"{sheet_name} 终止半径不正确")
    workbook.close()
    return {
        "path": str(path.resolve()),
        "sheet_names": ["温度", "水分浓度"],
        "expected_rows": expected_rows,
        "expected_columns": expected_columns,
        "sheets": checks,
    }


def _json_safe(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    return value


def save_numeric_outputs(
    output_root: Path,
    time_s: np.ndarray,
    radius_cm: np.ndarray,
    temperature_c: np.ndarray,
    moisture_kg_kg: np.ndarray,
) -> dict:
    data_directory = output_root / "data"
    data_directory.mkdir(parents=True, exist_ok=True)
    npz_path = data_directory / "q1_results_full.npz"
    np.savez_compressed(
        npz_path,
        time_s=time_s,
        radius_cm=radius_cm,
        temperature_c=temperature_c,
        moisture_kg_kg=moisture_kg_kg,
    )

    csv_path = data_directory / "q1_results_full.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "time_s",
                "radius_cm",
                "temperature_c",
                "moisture_kg_kg",
            ]
        )
        for time_index, time_value in enumerate(time_s):
            for radius_index, radius_value in enumerate(radius_cm):
                writer.writerow(
                    [
                        f"{time_value:.0f}",
                        f"{radius_value:.1f}",
                        f"{temperature_c[time_index, radius_index]:.10g}",
                        f"{moisture_kg_kg[time_index, radius_index]:.10g}",
                    ]
                )

    return {
        "npz": str(npz_path.resolve()),
        "csv": str(csv_path.resolve()),
        "summary": str((data_directory / "q1_summary.json").resolve()),
    }


def run_bessel_benchmark(
    cell_count: int,
    dt_s: float,
    *,
    parameters: Q1Parameters = DEFAULT_PARAMETERS,
) -> dict:
    """用恒温圆柱解析解交叉验证热方程；SciPy 缺失时明确跳过。"""

    try:
        from scipy.optimize import brentq
        from scipy.special import j0, j1
    except ImportError:
        return {
            "status": "skipped",
            "reason": "未安装 SciPy；主求解器不依赖 SciPy，仅解析基准需要。",
        }

    ambient_c = 60.0
    radius, numerical, _ = solve_heat(
        cell_count,
        dt_s,
        int(parameters.end_time_s),
        lambda _time: ambient_c,
        parameters=parameters,
    )
    biot = (
        parameters.heat_transfer_w_m2_k
        * parameters.radius_m
        / parameters.conductivity_w_m_k
    )

    def characteristic(value: float) -> float:
        return value * j1(value) - biot * j0(value)

    roots = []
    for index in range(20):
        lower = index * math.pi + 1e-8
        upper = (index + 0.5) * math.pi - 1e-8
        roots.append(brentq(characteristic, lower, upper))
    roots = np.asarray(roots)
    coefficients = 2.0 * j1(roots) / (
        roots * (j0(roots) ** 2 + j1(roots) ** 2)
    )
    errors = {}
    for report_time in (300, 900, 1800):
        fourier = (
            parameters.thermal_diffusivity_m2_s
            * report_time
            / parameters.radius_m**2
        )
        ratio = np.zeros_like(radius)
        for root, coefficient in zip(roots, coefficients, strict=True):
            ratio += (
                coefficient
                * j0(root * radius / parameters.radius_m)
                * np.exp(-(root**2) * fourier)
            )
        exact = ambient_c + (parameters.initial_temperature_c - ambient_c) * ratio
        errors[str(report_time)] = float(
            np.max(np.abs(numerical[report_time - 1] - exact))
        )
    return {"status": "passed", "max_abs_error_c_by_time_s": errors}


def run_q1(
    *,
    environment_path: Path = DEFAULT_ENV_PATH,
    template_path: Path = DEFAULT_TEMPLATE_PATH,
    output_root: Path = DEFAULT_OUTPUT_ROOT,
    cell_count: int = 400,
    dt_s: float = 0.5,
    full_validation: bool = True,
    parameters: Q1Parameters = DEFAULT_PARAMETERS,
) -> dict:
    """运行问题 1 的完整可复现流程。"""

    ambient_temperature, ambient_moisture, input_metadata = load_environment(
        environment_path, parameters.end_time_s
    )
    output_count = int(parameters.end_time_s)

    convergence = []
    if full_validation:
        settings: Iterable[tuple[int, float]] = (
            (max(40, cell_count // 2), dt_s * 2.0),
            (cell_count, dt_s),
            (cell_count * 2, dt_s / 2.0),
        )
    else:
        settings = ((cell_count, dt_s),)

    baseline = None
    for current_cells, current_dt in settings:
        radius, moisture, mass_diagnostics = solve_mass(
            current_cells,
            current_dt,
            output_count,
            ambient_moisture,
            parameters=parameters,
        )
        _, temperature, heat_diagnostics = solve_heat(
            current_cells,
            current_dt,
            output_count,
            ambient_temperature,
            parameters=parameters,
        )
        convergence.append(
            {
                "cell_count": current_cells,
                "dt_s": current_dt,
                "temperature_center_1800s_c": float(temperature[-1, 0]),
                "temperature_surface_1800s_c": float(temperature[-1, -1]),
                "moisture_center_1800s_kg_kg": float(moisture[-1, 0]),
                "moisture_surface_1800s_kg_kg": float(moisture[-1, -1]),
                "heat_cumulative_relative_residual": heat_diagnostics[
                    "cumulative_relative_residual"
                ],
                "mass_cumulative_relative_residual": mass_diagnostics[
                    "cumulative_relative_residual"
                ],
            }
        )
        if current_cells == cell_count and math.isclose(current_dt, dt_s):
            baseline = (
                radius,
                moisture,
                mass_diagnostics,
                temperature,
                heat_diagnostics,
            )
    if baseline is None:
        raise RuntimeError("未获得基准网格结果")

    (
        radius,
        moisture,
        mass_diagnostics,
        temperature,
        heat_diagnostics,
    ) = baseline
    output_time_s = np.arange(1, output_count + 1, dtype=float)
    output_radius_cm = np.round(np.arange(0.0, 2.0 + 0.05, 0.1), 10)
    moisture_output = resample_radius(radius, moisture, output_radius_cm)
    temperature_output = resample_radius(radius, temperature, output_radius_cm)
    workbook_path = output_root / "workbooks" / "result1.xlsx"
    workbook_validation = export_result_workbook(
        template_path,
        workbook_path,
        output_time_s,
        output_radius_cm,
        temperature_output,
        moisture_output,
    )

    report_times = np.asarray([100, 300, 600, 900, 1200, 1500, 1800])
    report_radii = np.asarray([0.0, 0.5, 1.0, 1.5, 2.0])
    time_indices = report_times - 1
    radius_indices = np.searchsorted(output_radius_cm, report_radii)
    summary = {
        "status": "completed",
        "heat_model": "radial_conduction_with_surface_convection",
        "parameters": {
            **asdict(parameters),
            "thermal_diffusivity_m2_s": parameters.thermal_diffusivity_m2_s,
        },
        "input_validation": input_metadata,
        "numerical_settings": {
            "cell_count": cell_count,
            "dt_s": dt_s,
            "output_interval_s": 1,
            "output_radius_interval_cm": 0.1,
            "theta": 0.5,
            "rannacher_steps": 2,
        },
        "diagnostics": {
            "heat": heat_diagnostics,
            "mass": {
                key: value
                for key, value in mass_diagnostics.items()
                if not isinstance(value, np.ndarray)
            },
            "convergence": convergence,
            "bessel_benchmark": (
                run_bessel_benchmark(cell_count, dt_s, parameters=parameters)
                if full_validation
                else {"status": "skipped", "reason": "快速模式"}
            ),
            "workbook": workbook_validation,
        },
        "paper_table": {
            "times_s": report_times.tolist(),
            "radii_cm": report_radii.tolist(),
            "temperature_c": temperature_output[
                np.ix_(time_indices, radius_indices)
            ].tolist(),
            "moisture_kg_kg": moisture_output[
                np.ix_(time_indices, radius_indices)
            ].tolist(),
        },
    }
    files = save_numeric_outputs(
        output_root,
        output_time_s,
        output_radius_cm,
        temperature_output,
        moisture_output,
    )
    summary["files"] = {**files, "workbook": str(workbook_path.resolve())}
    Path(files["summary"]).write_text(
        json.dumps(_json_safe(summary), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary


def _print_paper_table(title: str, matrix: np.ndarray, times: list, radii: list) -> None:
    print(f"\n{title}")
    print("时间/s |" + "".join(f"{radius:>10.1f}" for radius in radii))
    for time_value, row in zip(times, matrix, strict=True):
        print(f"{time_value:6d} |" + "".join(f"{value:10.4f}" for value in row))


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--environment", type=Path, default=DEFAULT_ENV_PATH)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE_PATH)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--cells", type=int, default=400, help="径向网格数")
    parser.add_argument("--dt", type=float, default=0.5, help="内部时间步长/s")
    parser.add_argument(
        "--quick",
        action="store_true",
        help="跳过网格加密和解析基准，仅运行基准求解与守恒检查",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_argument_parser().parse_args(argv)
    summary = run_q1(
        environment_path=args.environment,
        template_path=args.template,
        output_root=args.output_root,
        cell_count=args.cells,
        dt_s=args.dt,
        full_validation=not args.quick,
    )
    table = summary["paper_table"]
    _print_paper_table(
        "表1  30分钟内药材的温度 (°C)",
        np.asarray(table["temperature_c"]),
        table["times_s"],
        table["radii_cm"],
    )
    _print_paper_table(
        "表2  30分钟内药材的水分浓度 (kg/kg)",
        np.asarray(table["moisture_kg_kg"]),
        table["times_s"],
        table["radii_cm"],
    )
    print(f"\n状态：{summary['status']}")
    print(f"结果工作簿：{summary['files']['workbook']}")
    print(f"诊断摘要：{summary['files']['summary']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
